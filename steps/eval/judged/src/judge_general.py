"""LLM-judged general benchmarks (Granite: ProfBench, GDPval).

Both grade free-form work with an LLM judge reached through an OpenAI-compatible
endpoint (``judge_base_url`` / ``judge_model`` / ``judge_api_key_env``; default
Claude Sonnet 5 on the configured LiteLLM gateway). ``judge_model=self`` grades with the
served model instead (key-less smoke runs; ``details.judge_is_self``). Every
judge call's token usage (prompt, completion, cached) is kept with the example
it graded and totalled in ``details.judge_usage``.

ProfBench (metric "overall")
    NVlabs/ProfBench, pinned by commit: the report-generation protocol of the
    public leaderboard (lite version, 160 samples = 3-5 per task over 40 tasks,
    prompt only). The upstream judge call (``utils.get_criterion_fulfilment``:
    prompt, "mixed" reasoning effort, sampling) grades every rubric criterion,
    and upstream ``score_report_generation.get_predicted_score_per_task_id_e2e``
    turns the Yes/No ratings into the Overall score (weighted rubric score per
    response, mean per task, per domain, then over domains). ``value`` is
    Overall as a fraction.

GDPval (metric "Elo")
    openai/gdpval (gold subset, 220 tasks). The agent is Stirrup (the harness
    of Artificial Analysis' GDPval-AA) with a code-exec tool in a granite sandbox
    holding the task's reference files; its deliverable files are compared
    pairwise, in both orders, with the expert deliverable by the judge. GDPval-AA's
    judge prompt and reference submissions are not public, so its Elo cannot be
    reproduced here; the reference run judged with gemini-3.1-pro-preview,
    not this benchmark's default judge (Claude Sonnet 5 on the configured
    gateway). ``value`` is an Elo-style score derived from the win rate against the
    expert (expert = ``elo_anchor``), flagged as an approximation in results.json
    (``details.elo_is_approximation``, ``details.metric_note``) and in the log. It
    is relative (the expert sits at the anchor), not on the scale of
    published GDPval numbers.

Both judge one response per task (``--repeats 1``); ``details.pass_at_k`` is
the k = 1 record (GDPval: the win rate, no pass@k).

Per-example work goes under ``<output_dir>/samples/`` (ProfBench) or
``<output_dir>/tasks/`` (GDPval); finished examples are skipped on restart.
Both are splittable: ``--phase generate`` saves the policy's reports / the
agent's deliverables there without building the judge (no key needed),
``--phase score`` judges them without the served model.
"""

from __future__ import annotations

import asyncio
import base64
import collections
import concurrent.futures
import functools
import hashlib
import importlib.util
import io
import json
import logging
import math
import os
import re
import shlex
import sys
import threading
import types
import urllib.request
import zipfile
from pathlib import Path
from typing import Any, ClassVar

from granite_evals import data
from granite_evals.registry import Benchmark, failure_policy, pass_at_k_record, register

log = logging.getLogger(__name__)

DEFAULT_JUDGE_BASE_URL = ""  # set --option judge_base_url=<OpenAI-compatible judge endpoint>
DEFAULT_JUDGE_MODEL = "aws/claude-sonnet-5"
DEFAULT_JUDGE_KEY_ENV = "GRANITE_EVALS_JUDGE_API_KEY"
# USD per million tokens, for the cost estimate in details.judge_usage
# (Claude Sonnet 5: $2 in / $10 out, cache reads ~90% off, cache writes +25%).
PRICE_IN, PRICE_OUT, CACHE_READ_FACTOR, CACHE_WRITE_FACTOR = 2.0, 10.0, 0.1, 1.25

# Environment variables never handed to a sandbox (model-written code runs there).
_SECRET_ENV = re.compile(r"(KEY|TOKEN|SECRET|PASSWORD|CREDENTIAL)", re.I)


def _count(values) -> dict[str, int]:
    return dict(collections.Counter(values))


def _write_json(path: Path, obj: Any) -> None:
    tmp = path.with_suffix(path.suffix + ".partial")
    tmp.write_text(json.dumps(obj, indent=2, default=str))
    tmp.replace(path)


# ---------------------------------------------------------------------------
# Judge client
# ---------------------------------------------------------------------------


def usage_record(usage: Any) -> dict[str, int]:
    """prompt / completion / cached (read) / cache-write tokens of one call.

    OpenAI-style ``prompt_tokens_details.cached_tokens`` and Anthropic-style
    ``cache_read_input_tokens`` / ``cache_creation_input_tokens`` (LiteLLM
    passes both through) are both understood. ``prompt_tokens`` includes the
    cached tokens.
    """
    if usage is None:
        return {"prompt_tokens": 0, "completion_tokens": 0, "cached_tokens": 0, "cache_write_tokens": 0,
                "reasoning_tokens": 0}
    u = usage.model_dump() if hasattr(usage, "model_dump") else dict(usage)
    details = u.get("prompt_tokens_details") or {}
    cached = details.get("cached_tokens") or u.get("cache_read_input_tokens") or 0
    return {
        "prompt_tokens": int(u.get("prompt_tokens") or 0),
        "completion_tokens": int(u.get("completion_tokens") or 0),
        "cached_tokens": int(cached),
        "cache_write_tokens": int(u.get("cache_creation_input_tokens") or 0),
        # Thinking tokens (part of completion_tokens), when the endpoint reports them.
        "reasoning_tokens": int((u.get("completion_tokens_details") or {}).get("reasoning_tokens") or 0),
    }


def total_usage(records: list[dict], n_examples: int) -> dict[str, Any]:
    """Sum per-call usage records, with per-example averages and a cost estimate."""
    keys = ("prompt_tokens", "completion_tokens", "cached_tokens", "cache_write_tokens", "reasoning_tokens")
    tot = {k: sum(r.get(k, 0) for r in records) for k in keys}
    uncached = max(tot["prompt_tokens"] - tot["cached_tokens"] - tot["cache_write_tokens"], 0)
    cost = (
        uncached * PRICE_IN
        + tot["cached_tokens"] * PRICE_IN * CACHE_READ_FACTOR
        + tot["cache_write_tokens"] * PRICE_IN * CACHE_WRITE_FACTOR
        + tot["completion_tokens"] * PRICE_OUT
    ) / 1e6
    return {
        "calls": len(records),
        **tot,
        "examples": n_examples,
        "per_example": {k: round(v / n_examples, 1) for k, v in tot.items()} if n_examples else {},
        "calls_per_example": round(len(records) / n_examples, 2) if n_examples else 0,
        "estimated_cost_usd": round(cost, 4),
        "cost_per_example_usd": round(cost / n_examples, 5) if n_examples else 0,
        "price_assumption": f"${PRICE_IN}/M input, ${PRICE_OUT}/M output, cache read x{CACHE_READ_FACTOR}, "
        f"cache write x{CACHE_WRITE_FACTOR} (Claude Sonnet 5 list prices)",
    }


class Judge:
    """OpenAI-compatible judge endpoint with the ``client.chat.completions.create``
    shape the harnesses call, adding:

    - the judge model name (the harness's own ``model`` argument is replaced),
    - prompt caching: a ``cache_split(text) -> (prefix, rest)`` hook turns the
      user message into two text blocks, the first marked ``cache_control``
      (Anthropic prompt caching through LiteLLM); the concatenated text is
      unchanged. Off for ``judge_model=self``.
    - request parameters the endpoint rejects with a 400 (e.g. temperature and
      top_p together on newer Claude models) are dropped once, recorded in
      ``adaptations``, and not sent again,
    - ``extra_body``: provider fields merged into every request body (e.g.
      Anthropic ``thinking`` / ``output_config`` through LiteLLM); never dropped,
    - the usage of the last call on this thread (``last_usage``).
    """

    library = "openrouter"  # ProfBench's marker for the chat.completions path

    def __init__(
        self,
        *,
        base_url: str,
        model: str,
        api_key: str,
        is_self: bool,
        cache_split=None,
        extra_body: dict | None = None,
        timeout: float = 600,
        max_retries: int = 5,
        client: Any = None,
    ):
        if client is None:
            from openai import OpenAI

            client = OpenAI(base_url=base_url, api_key=api_key or "EMPTY", timeout=timeout, max_retries=max_retries)
        self._client = client
        self.base_url, self.model, self.is_self = base_url, model, is_self
        self.cache_split = None if is_self else cache_split
        self.extra_body = dict(extra_body or {})
        self.dropped: list[str] = []
        self.adaptations: list[str] = []
        self._lock = threading.Lock()
        self._local = threading.local()
        self.chat = types.SimpleNamespace(completions=types.SimpleNamespace(create=self.create))

    def describe(self) -> dict[str, Any]:
        return {
            "judge_model": self.model,
            "judge_base_url": self.base_url,
            "judge_is_self": self.is_self,
            "judge_prompt_caching": self.cache_split is not None,
            "judge_extra_body": dict(self.extra_body),
            "judge_dropped_params": list(self.dropped),
            "judge_adaptations": list(self.adaptations),
        }

    def last_usage(self) -> dict[str, int]:
        return getattr(self._local, "usage", usage_record(None))

    def _cached_messages(self, messages: list[dict]) -> list[dict]:
        if self.cache_split is None:
            return messages
        out = []
        for m in messages:
            if m.get("role") == "user" and isinstance(m.get("content"), str):
                parts = self.cache_split(m["content"])
                if parts and parts[0] and parts[1]:
                    m = {
                        **m,
                        "content": [
                            {"type": "text", "text": parts[0], "cache_control": {"type": "ephemeral"}},
                            {"type": "text", "text": parts[1]},
                        ],
                    }
            out.append(m)
        return out

    def create(self, **request):
        request = {**request, "model": self.model}
        request["messages"] = self._cached_messages(request["messages"])
        if self.extra_body:
            request["extra_body"] = {**self.extra_body, **(request.get("extra_body") or {})}
        with self._lock:
            for key in self.dropped:
                request.pop(key, None)
        while True:
            try:
                completion = self._client.chat.completions.create(**request)
                break
            except Exception as e:
                if getattr(e, "status_code", None) != 400:
                    raise
                message = str(e).lower()
                if "cache_control" in message and self.cache_split is not None:
                    with self._lock:
                        self.cache_split = None
                        self.adaptations.append("prompt caching disabled (endpoint rejected cache_control)")
                    request["messages"] = [
                        {**m, "content": "".join(b["text"] for b in m["content"])} if isinstance(m["content"], list) else m
                        for m in request["messages"]
                    ]
                    continue
                # Drop the first optional parameter the error names. Any other 400 (context
                # length, a bad extra_body field, ...) is raised: dropping a parameter the
                # endpoint did not object to would change the judge without fixing anything.
                optional = [k for k in ("top_p", "temperature", "reasoning_effort") if k in request]
                named = [k for k in optional if k.replace("_", "") in message.replace("_", "")]
                if not named:
                    raise
                key = named[0]
                request.pop(key)
                with self._lock:
                    if key not in self.dropped:
                        self.dropped.append(key)
                        self.adaptations.append(f"dropped {key}: {str(e)[:300]}")
                log.warning("judge rejected %s; retrying without it", key)
        self._local.usage = usage_record(getattr(completion, "usage", None))
        return completion


class JudgedBenchmark(Benchmark):
    extra = "judged"

    def opt(self, key: str, default: Any) -> Any:
        """A ``--option key=value`` (always a string on the CLI), cast like ``default``."""
        value = self.config.options.get(key)
        if value is None:
            return default
        if isinstance(default, bool):
            return str(value).lower() in ("1", "true", "yes", "on")
        return type(default)(value)

    @property
    def judge_model(self) -> str:
        return self.opt("judge_model", DEFAULT_JUDGE_MODEL)

    def make_judge(self, base_url: str, served: str, cache_split=None) -> Judge:
        if self.judge_model == "self":
            return Judge(base_url=base_url, model=served, api_key="EMPTY", is_self=True)
        key_env = self.opt("judge_api_key_env", DEFAULT_JUDGE_KEY_ENV)
        key = os.environ.get(key_env, "")
        if not key:
            raise SystemExit(
                f"{self.id}: judge {self.judge_model!r} needs an API key in ${key_env} "
                "(or use --option judge_model=self for a key-less smoke run)"
            )
        from granite_evals import meter

        # The CLI already routes a --option judge_base_url through its meter; the class
        # default is a paid endpoint too and must be metered here (never twice).
        if not (self.config.options.get("judge_base_url") or DEFAULT_JUDGE_BASE_URL):
            raise SystemExit(f"{self.id}: no judge endpoint; pass --option judge_base_url=<OpenAI-compatible URL> (or judge_model=self)")
        base = self.config.options.get("judge_base_url") or meter.metered(DEFAULT_JUDGE_BASE_URL, role="judge")
        return Judge(
            base_url=base,
            model=self.judge_model,
            api_key=key,
            is_self=False,
            cache_split=cache_split if self.opt("judge_cache", True) else None,
            extra_body=self.judge_extra_body(),
            timeout=self.opt("judge_timeout", 600.0),
        )

    def judge_extra_body(self) -> dict[str, Any]:
        """Provider request fields for the judge, from the options

        judge_thinking=adaptive|enabled:<budget_tokens>|disabled
                          Anthropic ``thinking`` (LiteLLM passes it through)
        judge_effort=low|medium|high|xhigh|max
                          Anthropic ``output_config.effort``
        judge_extra_body=<json object>
                          anything else, merged last

        (``judge_reasoning_effort`` is the OpenAI-style top-level field instead.)
        """
        body: dict[str, Any] = {}
        thinking = self.opt("judge_thinking", "")
        if thinking in ("adaptive", "disabled"):
            body["thinking"] = {"type": thinking}
        elif thinking.startswith("enabled:"):
            body["thinking"] = {"type": "enabled", "budget_tokens": int(thinking.split(":", 1)[1])}
        elif thinking:
            raise SystemExit(f"{self.id}: judge_thinking must be adaptive, enabled:<budget_tokens> or disabled")
        if effort := self.opt("judge_effort", ""):
            body["output_config"] = {"effort": effort}
        if extra := self.opt("judge_extra_body", ""):
            parsed = json.loads(extra)
            if not isinstance(parsed, dict):
                raise SystemExit(f"{self.id}: judge_extra_body must be a JSON object")
            body.update(parsed)
        return body

    def one_sample(self) -> None:
        """Both judged benchmarks grade one response per task (``--repeats 1``):
        their scores (ProfBench's weighted rubric Overall, GDPval's win rate)
        are defined over one response, and k responses would cost k paid judge
        passes for a number upstream does not report."""
        if self.repeats != 1:
            raise SystemExit(f"{self.id}: one response per task is judged; use --repeats 1")

    def sampling(self) -> dict[str, Any]:
        """Policy-model sampling overrides; unset ones fall back to the
        checkpoint's generation_config, which vLLM applies."""
        out: dict[str, Any] = {}
        for key, cast in (("temperature", float), ("top_p", float)):
            if key in self.config.options:
                out[key] = cast(self.config.options[key])
        return out


# ---------------------------------------------------------------------------
# ProfBench
# ---------------------------------------------------------------------------

PROFBENCH_REPO = "NVlabs/ProfBench"
PROFBENCH_COMMIT = "b06a29cda4d1433e9a9aad8171e4086299083e94"
PROFBENCH_FILES = {
    "utils.py": "6aa796c531e1601369cd7249e1f33b0a941e993ef90a7e9004170ee01665d3fe",
    "score_report_generation.py": "f8cf57e2b02ba4e667281e27c942bd5d5b9e8ea3550d51ff565a71bc8b4ac179",
    "run_report_generation.py": "e24beb96496250a4f372d5626bd94681aabca870538e9708cf8086727ca9d55f",
    "score_llm_judge.py": "d9c925aa95402162cba1dcca200822ec43b0ddc23c8e25a2e29f87649ccfe059",
    "LICENSE": "5dc788c299ee9525feb4053830c1358c529c67a906e461311e2200a89f0d5ffc",
}
PROFBENCH_RESPONSE_MODELS = ("o3", "grok4", "r1-0528")
# Upstream's judge prompt is "Response:\n\n{response}\n\nEvaluate whether ...": the
# response is the prefix shared by every criterion of one sample.
_PB_SPLIT = "\n\nEvaluate whether the response above satisfies this criterion: "


def _pb_cache_split(text: str) -> tuple[str, str] | None:
    i = text.rfind(_PB_SPLIT)
    return (text[:i], text[i:]) if i > 0 else None


def profbench_dir() -> Path:
    """The upstream ProfBench scripts at the pinned commit, sha256-verified.

    The judged image bakes them into /opt/profbench/<commit> (docker/extras/judged.sh);
    elsewhere they are fetched once into ~/.cache/granite/profbench/<commit>.
    """
    candidates = [os.environ.get("GRANITE_EVALS_PROFBENCH_DIR", ""), f"/opt/profbench/{PROFBENCH_COMMIT}"]
    for c in filter(None, candidates):
        if all((Path(c) / name).exists() for name in PROFBENCH_FILES):
            return Path(c)
    d = Path.home() / ".cache" / "granite" / "profbench" / PROFBENCH_COMMIT
    d.mkdir(parents=True, exist_ok=True)
    for name, sha in PROFBENCH_FILES.items():
        path = d / name
        if path.exists() and hashlib.sha256(path.read_bytes()).hexdigest() == sha:
            continue
        url = f"https://raw.githubusercontent.com/{PROFBENCH_REPO}/{PROFBENCH_COMMIT}/{name}"
        with urllib.request.urlopen(url, timeout=60) as r:
            content = r.read()
        if hashlib.sha256(content).hexdigest() != sha:
            raise RuntimeError(f"{url}: sha256 mismatch")
        path.write_bytes(content)
    return d


@functools.cache
def profbench_modules() -> types.SimpleNamespace:
    """Import the upstream scripts (they are plain modules, not a package)."""
    d = profbench_dir()

    def load(name: str) -> types.ModuleType:
        spec = importlib.util.spec_from_file_location(f"profbench_{name}", d / f"{name}.py")
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        return module

    utils = load("utils")
    # run_report_generation does `from utils import ...`.
    saved = sys.modules.get("utils")
    sys.modules["utils"] = utils
    try:
        generation = load("run_report_generation")
    finally:
        if saved is None:
            sys.modules.pop("utils", None)
        else:
            sys.modules["utils"] = saved
    for m in (utils, generation):
        m.print = lambda *a, **k: None  # upstream progress prints
    return types.SimpleNamespace(
        utils=utils,
        generation=generation,
        score=load("score_report_generation"),
        judge_score=load("score_llm_judge"),
    )


@register
class ProfBench(JudgedBenchmark):
    """ProfBench report generation (lite), graded by rubric with an LLM judge.

    Options:
      version=lite|full|debug   upstream sample sets (lite = the leaderboard's 160)
      responses=model|o3|grok4|r1-0528
                                grade the served model's reports (default) or the
                                reports shipped with the dataset
      judge_model=<name>|self|human
                                ``human`` uses the dataset's human fulfilment labels
                                (only with responses=o3|grok4|r1-0528; no model, no key)
      max_tokens=64000          policy generation cap (upstream's reasoning setting)
      temperature, top_p        policy sampling (default: generation_config)
      judge_workers=16          concurrent samples being judged
      judge_reasoning_effort=   one effort for every criterion, sent by upstream's
                                call as ``reasoning_effort`` (unset = upstream's
                                mixed high/low)
      judge_thinking, judge_effort, judge_extra_body
                                provider thinking fields (JudgedBenchmark.judge_extra_body)
      judge_tag=                keep this judge's ratings apart (judgments-<tag>.json),
                                e.g. to re-judge the same reports with other settings
      judge_max_criteria=0      probes: judge at most this many criteria per sample
                                (the rest are left unjudged and not scored)
      max_failed_frac=0.05      see below

    A failed generation drops its sample and a failed judge call drops its
    criterion, as upstream does; both are retried on resume. ``--phase score``
    judges the saved ``response.json`` files; a sample without one counts as a
    failed generation (status ``not_generated``). ``details.criteria_failed``
    / ``criteria_total`` count the criteria left unrated; above ``max_failed_frac``
    (0 = any) the run fails, below it the result is flagged ``incomplete``.

    The verdict is the judge message's ``content`` only (upstream's
    ``startswith("Yes")``); reasoning / thinking fields of the response are never
    read. Ratings that start with neither Yes nor No are counted in
    ``details.judge_unparsed_ratings``.
    """

    id = "profbench"
    metric = "overall"
    harness_packages = ("openai", "numpy", "scikit-learn")
    dataset = "nvidia/ProfBench"
    dataset_revision = "f5c471cf2f277c264621e4b5c0cad5bb466f9260"
    split: ClassVar[str] = "test"
    splittable = True

    @property
    def responses(self) -> str:
        r = self.opt("responses", "model")
        if r not in ("model", *PROFBENCH_RESPONSE_MODELS):
            raise SystemExit(f"profbench: responses must be model or one of {PROFBENCH_RESPONSE_MODELS}")
        return r

    def needs_server(self) -> bool:
        return self.responses == "model"

    def score_needs_server(self) -> bool:
        return self.judge_model == "self"

    # -- entry point -------------------------------------------------------

    def run(self, base_url: str, served_model_name: str) -> dict[str, Any]:
        self.one_sample()
        if self.judge_model == "human" and self.responses == "model":
            raise SystemExit("profbench: judge_model=human needs responses=o3|grok4|r1-0528 (human labels)")
        pb = profbench_modules()
        source, revision = self.dataset_source()
        rows = data.load_split(source, revision=revision, split=self.split)
        tasks = data.take(rows, self.config.limit, key="task_id")
        samples = self._samples(pb, rows, tasks)
        log.info("profbench: %d tasks, %d samples (responses=%s, judge=%s)",
                 len(tasks), len(samples), self.responses, self.judge_model)

        root = self.config.output_dir / "samples"

        with concurrent.futures.ThreadPoolExecutor(max_workers=self.config.workers) as pool:
            gens = list(pool.map(lambda s: self._generate(s, root, base_url, served_model_name), samples))
        if not self.scoring:
            return self._generated(samples, gens, source, revision)
        judge = None if self.judge_model == "human" else self.make_judge(base_url, served_model_name, _pb_cache_split)
        with concurrent.futures.ThreadPoolExecutor(max_workers=self.opt("judge_workers", 16)) as pool:
            judged = list(pool.map(lambda sg: self._judge_sample(pb, sg[0], sg[1], root, judge), zip(samples, gens)))
        return self._aggregate(pb, samples, gens, judged, judge, source, revision)

    def _samples(self, pb, rows: list[dict], tasks: list[dict]) -> list[dict]:
        if self.responses != "model":
            return [{"task": t, "k": 0} for t in tasks]
        version = self.opt("version", "lite")
        if version == "debug":
            # Upstream's debug set is the dataset's first row; here: one sample per selected task.
            counts = collections.Counter(t["task_id"] for t in tasks)
        else:
            # Upstream's sample counts per task for this version (it needs every task id).
            expanded = pb.generation.filter_data([{"task_id": r["task_id"]} for r in rows], version)
            counts = collections.Counter(dp["task_id"] for dp in expanded)
        cap = self.opt("samples", 0)  # cheap checks: at most this many samples per task
        return [{"task": t, "k": k} for t in tasks
                for k in range(min(counts.get(t["task_id"], 0), cap or 10**9))]

    # -- generation ----------------------------------------------------------

    def _generate(self, sample: dict, root: Path, base_url: str, served: str) -> dict:
        task, k = sample["task"], sample["k"]
        sdir = root / task["task_id"] / str(k)
        path = sdir / "response.json"
        if path.exists():
            return json.loads(path.read_text())
        sdir.mkdir(parents=True, exist_ok=True)
        if self.responses != "model":
            # The dataset's reports: no model, so any phase may write them.
            out = {"status": "dataset", "response": task[f"{self.responses}_response"],
                   "prompt_tokens": None, "completion_tokens": None}
            _write_json(path, out)
            return out
        if not self.generating:
            # Score phase: never generate; the sample is left out like a failed generation.
            log.warning("profbench: %s", self.not_generated(f"{task['task_id']}#{k}"))
            return {"status": "not_generated", "response": None, "prompt_tokens": None, "completion_tokens": None}
        try:
            from openai import OpenAI

            client = OpenAI(base_url=base_url, api_key="EMPTY", timeout=self.opt("timeout", 3600.0))
            request = {
                # Upstream's chat.completions request: the prompt as the only user message.
                "model": served,
                "messages": [{"role": "user", "content": task["prompt"]}],
                "max_tokens": self.opt("max_tokens", 64000),
                "seed": self.config.seed + k,
                **self.sampling(),
            }
            completion = client.chat.completions.create(**request)
            choice = completion.choices[0]
            content = choice.message.content or ""
            out = {
                "status": "ok" if content.strip() else "empty_response",
                "finish_reason": choice.finish_reason,
                "response": content,
                "reasoning": getattr(choice.message, "reasoning_content", None) or getattr(choice.message, "reasoning", None),
                "prompt_tokens": completion.usage.prompt_tokens,
                "completion_tokens": completion.usage.completion_tokens,
            }
        except Exception as e:  # one broken sample must not sink the run
            log.exception("profbench: generation %s#%d failed", task["task_id"], k)
            # Not persisted: retried on resume.
            return {"status": f"error:{type(e).__name__}", "response": None, "prompt_tokens": None, "completion_tokens": None}
        _write_json(path, out)
        log.info("profbench: generated %s#%d %s (%s tokens)", task["task_id"], k, out["status"], out["completion_tokens"])
        return out

    def _generated(self, samples, gens, source, revision) -> dict[str, Any]:
        """The generate phase's summary (generation.json): reports saved, not judged."""
        failed = sum(g["response"] is None for g in gens)
        if failed:
            log.warning("profbench: %d/%d generations failed; rerun --phase generate to retry them",
                        failed, len(gens))
        return {
            "n": len(gens) - failed,  # samples with a saved response.json
            "samples": len(samples),
            "generations_failed": failed,
            "statuses": _count(g["status"] for g in gens),
            "dataset": source,
            "dataset_revision": revision,
            "version": self.opt("version", "lite") if self.responses == "model" else "provided-reports",
            "responses": self.responses,
            "tasks": sorted({s["task"]["task_id"] for s in samples}),
            "sampling": {**self.sampling(), "max_tokens": self.opt("max_tokens", 64000),
                         "defaults": "generation_config of the checkpoint"},
        }

    # -- judging -------------------------------------------------------------

    def _criteria(self, task: dict, response: str) -> list[dict]:
        """Upstream's (response x criterion) datapoints (run_best_llm_judge_on_generated_reports.load_data)."""
        return [
            {
                "task_id": task["task_id"],
                "domain": task["domain"],
                "criterion_description": c["criterion_description"],
                "criterion_weight": c["criterion_weight"],
                "criterion_type": c["criterion_type"],
                "response": response,
            }
            for c in task["rubrics"]
        ]

    def _judge_sample(self, pb, sample: dict, gen: dict, root: Path, judge: Judge | None) -> dict:
        """Ratings for every criterion of one sample: {idx: {judge_rating, usage}}."""
        task, k = sample["task"], sample["k"]
        tag = self.opt("judge_tag", "")
        path = root / task["task_id"] / str(k) / (f"judgments-{tag}.json" if tag else "judgments.json")
        done: dict[str, dict] = json.loads(path.read_text()) if path.exists() else {}
        if gen["response"] is None:
            return {"status": "not_generated" if gen["status"] == "not_generated" else "generation_failed",
                    "ratings": {}}
        points = self._criteria(task, gen["response"])
        if self.judge_model == "human":
            for i, c in enumerate(task["rubrics"]):
                done[str(i)] = {"judge_rating": "Yes" if c[f"{self.responses}_fulfilment"] else "No", "usage": None}
        elif not gen["response"].strip():
            # Nothing to grade: every criterion unmet, no judge calls.
            for i in range(len(points)):
                done.setdefault(str(i), {"judge_rating": "No", "usage": None, "skipped": "empty_response"})
        else:
            todo = [i for i in range(len(points)) if str(i) not in done]
            if cap := self.opt("judge_max_criteria", 0):
                todo = [i for i in todo if i < cap]
            # None = upstream's mixed effort; a string is sent as reasoning_effort on every call.
            hyper = {"reasoning": self.opt("judge_reasoning_effort", "") or None}

            def one(i: int) -> None:
                try:
                    dp = pb.utils.get_criterion_fulfilment(dict(points[i]), i, hyper, judge, judge.model)
                    done[str(i)] = {"judge_rating": dp["judge_rating"] or "", "usage": judge.last_usage()}
                except Exception as e:
                    log.warning("profbench: judge %s#%d criterion %d failed: %s", task["task_id"], k, i, e)

            # First criterion alone so the response prefix is cached before the rest read it.
            if todo:
                one(todo[0])
            with concurrent.futures.ThreadPoolExecutor(max_workers=self.opt("judge_parallel", 4)) as inner:
                list(inner.map(one, todo[1:]))
        _write_json(path, done)
        missing = len(points) - len(done)
        if missing:
            log.warning("profbench: %s#%d: %d criteria unjudged (retried on resume)", task["task_id"], k, missing)
        log.info("profbench: judged %s#%d %d/%d criteria met", task["task_id"], k,
                 sum(str(d["judge_rating"]).startswith("Yes") for d in done.values()), len(points))
        return {"status": "judged" if not missing else "judge_incomplete", "ratings": done}

    # -- scoring -------------------------------------------------------------

    def _expected_criteria(self, task: dict, gen: dict) -> int:
        """Criteria a sample should have ratings for: all of them, or the first
        judge_max_criteria of a response the judge grades."""
        n = len(task["rubrics"])
        cap = self.opt("judge_max_criteria", 0)
        graded_by_judge = self.judge_model != "human" and (gen["response"] is None or gen["response"].strip())
        return min(n, cap) if cap and graded_by_judge else n

    def _aggregate(self, pb, samples, gens, judged, judge, source, revision) -> dict[str, Any]:
        rows, usage, statuses = [], [], []
        failed = total = scored = 0
        for sample, gen, j in zip(samples, gens, judged):
            task = sample["task"]
            statuses.append(gen["status"] if j["status"] == "judged" else j["status"])
            expected = self._expected_criteria(task, gen)
            total += expected
            if gen["response"] is None:
                # Generation failed: the sample is dropped, as upstream's parallel_launcher
                # drops an item whose worker raised (utils.py:44-59); counted as failed.
                failed += expected
                continue
            ratings = j["ratings"]
            failed += sum(str(i) not in ratings for i in range(expected))
            scored += any(str(i) in ratings for i in range(len(task["rubrics"])))
            for i, dp in enumerate(self._criteria(task, gen["response"])):
                r = ratings.get(str(i))
                if r is None:  # judge failed on this criterion: dropped, as upstream does
                    continue
                if r.get("usage"):
                    usage.append(r["usage"])
                rows.append({
                    **dp,
                    "prompt_tokens": gen["prompt_tokens"],
                    "completion_tokens": gen["completion_tokens"],
                    "judge_rating": r["judge_rating"],
                    "human_annotation": task["rubrics"][i].get(f"{self.responses}_fulfilment"),
                })
        policy = failure_policy(self.id, "criteria", failed, total, self.opt("max_failed_frac", 0.05))
        if failed:
            log.warning("profbench: %d/%d criteria not rated (generation or judge failed), left out; "
                        "rerun to retry them", failed, total)  # fmt: skip
        if not rows:
            raise SystemExit("profbench: nothing was judged")
        # The upstream scorer averages token counts over tasks; guard all-None fields.
        scoring_rows = [
            {**r, "prompt_tokens": r["prompt_tokens"] or 0, "completion_tokens": r["completion_tokens"] or 0}
            for r in rows
        ]
        if not any(r["response"] for r in scoring_rows):
            # Upstream's response-length stat is the mean of an empty list then (a crash).
            scoring_rows = [{**r, "response": " "} for r in scoring_rows]
        scores = pb.score.get_predicted_score_per_task_id_e2e(condition="judge_rating", value="Yes", data=scoring_rows)
        scores = {k: (v.item() if hasattr(v, "item") else v) for k, v in scores.items()}
        out: dict[str, Any] = {
            "value": scores["Overall"] / 100,
            "n": scored,  # samples with at least one rated criterion
            **policy,
            "dataset": source,
            "dataset_revision": revision,
            "harness": f"github.com/{PROFBENCH_REPO}@{PROFBENCH_COMMIT}",
            "version": self.opt("version", "lite") if self.responses == "model" else "provided-reports",
            "responses": self.responses,
            "scores": scores,
            "pass_at_k": pass_at_k_record(1, scores["Overall"] / 100, scores["Overall"] / 100, scored,
                                          "ProfBench Overall, one report per task"),  # fmt: skip
            "tasks": sorted({s["task"]["task_id"] for s in samples}),
            "statuses": _count(statuses),
            "criteria_judged": len(rows),
            "sampling": {**self.sampling(), "max_tokens": self.opt("max_tokens", 64000),
                         "defaults": "generation_config of the checkpoint"},
        }
        if judge is not None:
            out.update(judge.describe())
            effort = self.opt("judge_reasoning_effort", "")
            out["judge_request"] = (
                "upstream utils.get_criterion_fulfilment: "
                + (f"reasoning_effort {effort} for every criterion" if effort else
                   "reasoning mixed (high for Physics/Chemistry PhD or Style criteria, else low)")
                + ", temperature 0.6, top_p 0.95, max_tokens 32768 (minus judge_dropped_params)"
                + (f", extra body {json.dumps(judge.extra_body, sort_keys=True)}" if judge.extra_body else "")
            )
            out["judge_reasoning_effort"] = effort or "mixed"
            out["judge_tag"] = self.opt("judge_tag", "")
            if cap := self.opt("judge_max_criteria", 0):
                out["judge_max_criteria"] = cap
            out["judge_unparsed_ratings"] = sum(
                not str(r["judge_rating"]).startswith(("Yes", "No")) for r in rows)
            out["judge_usage"] = total_usage(usage, scored)
        else:
            out["judge_model"] = "human"
            out["judge_is_self"] = False
        if self.responses != "model" and self.judge_model != "human":
            out["judge_agreement"] = self._agreement(pb, rows)
        log.info("profbench: Overall %.1f over %d samples (%s)", scores["Overall"], scored,
                 ", ".join(f"{k} {v}" for k, v in scores.items() if k.endswith(("PhD", "MBA"))))
        return out

    @staticmethod
    def _agreement(pb, rows: list[dict]) -> dict[str, Any]:
        """Judge vs. human labels on the provided reports: upstream's macro-F1 (score_llm_judge)."""
        m = pb.judge_score
        m.data = [{**r, "judge_rating": str(r["judge_rating"])} for r in rows]
        out = {"macro_f1": m.get_metric()}
        for field, values in (("domain", ["Physics PhD", "Chemistry PhD", "Finance MBA", "Consulting MBA"]),
                              ("criterion_type", ["Extraction (recall)", "Reasoning", "Style"])):
            for v in values:
                if any(v in r[field] for r in rows):
                    out[v] = m.get_metric(condition=field, value=v)
        out["mean_pred_minus_human"] = m.get_mean_pred_error()
        return {k: (v.item() if hasattr(v, "item") else v) for k, v in out.items()}


# ---------------------------------------------------------------------------
# GDPval
# ---------------------------------------------------------------------------

GDPVAL_METRIC_NOTE = (
    "APPROXIMATION, not GDPval-AA Elo: GDPval-AA's judge prompt and reference "
    "submissions are not public (the reference run judged with "
    "gemini-3.1-pro-preview, not this benchmark's default judge). This is an "
    "Elo-style score from the pairwise win rate "
    "(ties = 0.5, both presentation orders) of the model's deliverables against the expert "
    "deliverables of the GDPval gold set, judged by {judge}: elo = anchor + 400*log10(p/(1-p)) "
    "with the expert at anchor={anchor} and p smoothed as (score+0.5)/(n+1). Not comparable "
    "with Artificial Analysis' published GDPval-AA numbers, nor on the scale of published "
    "Granite GDPval numbers: it is relative to the expert, and only runs with the same judge compare."
)

# The deliverable guidance of the GDPval paper's agent prompt (arXiv 2510.04374, A.3),
# without the steps that need an image viewer, which this harness does not offer.
GDPVAL_INSTRUCTIONS = """\
You are completing a real work task for a client. Deliverables:
- Any deliverables the user asked for must be files you create in the working directory, NOT only text in your final message. Report every deliverable file in the finish tool's paths.
- If the prompt does not request a specific filetype, use standard filetypes like PDF, PPTX, DOCX, XLSX, etc. Python (python-docx, openpyxl, python-pptx, reportlab, pandas, matplotlib) is available; install anything else you need.
- Your submission should be complete and self-contained. Even if you cannot fully complete the task due to limitations of the environment, produce as close to a complete solution as possible, and say in the finish reason (4 sentences max) what could not be done.
- Always be clear and comprehensive, but avoid extra verbosity.
- Avoid emojis, nonstandard bullet points and other special characters. Make plots and diagrams legible.
- Before finishing, open each deliverable file programmatically to check that it is not corrupted and that it contains exactly what the task asked for."""

GDPVAL_JUDGE_PROMPT = """\
You are an experienced {occupation} ({sector}) grading work submitted for a real task in your field. \
Two submissions, A and B, were produced independently for the task below. Decide which one a client \
would rather receive, judging correctness, completeness, adherence to the instructions, and professional \
quality and formatting. The grading rubric written by the task author is guidance on what matters.

# Task
{prompt}

# Reference files provided with the task
{references}

# Grading rubric (guidance)
{rubric}

Files are shown as extracted text; binary content that cannot be shown (images, audio, video) is listed \
by name and size only. Do not prefer a submission for its length or its position.

=== SUBMISSION A ===
{a}
=== END OF SUBMISSION A ===

=== SUBMISSION B ===
{b}
=== END OF SUBMISSION B ===

Compare the two submissions, then end your answer with exactly one line: "VERDICT: A", "VERDICT: B" or "VERDICT: TIE"."""

_GDP_SPLIT = "=== SUBMISSION A ==="
_VERDICT = re.compile(r"VERDICT\s*[:\-]?\s*\**\s*\[?\s*(A|B|TIE)\b", re.I)
TEXT_SUFFIXES = {".txt", ".md", ".csv", ".tsv", ".json", ".yaml", ".yml", ".py", ".html", ".htm", ".xml",
                 ".tex", ".sql", ".js", ".ts", ".r", ".ipynb", ".overpassql", ".log", ".ini", ".cfg"}


def _gdp_cache_split(text: str) -> tuple[str, str] | None:
    i = text.find(_GDP_SPLIT)
    return (text[:i], text[i:]) if i > 0 else None


def parse_verdict(text: str | None) -> str | None:
    matches = _VERDICT.findall(text or "")
    return matches[-1].upper() if matches else None


def elo_style(score: float, n: int, anchor: float) -> float:
    """Elo-style rating from a smoothed win rate against the anchor (the expert)."""
    p = (score + 0.5) / (n + 1)
    return anchor + 400 * math.log10(p / (1 - p))


def render_file(name: str, content: bytes, limit: int = 40_000) -> str:
    """A deliverable as text for the judge."""
    suffix = Path(name).suffix.lower()
    try:
        if suffix == ".docx":
            import docx

            d = docx.Document(io.BytesIO(content))
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                parts += [" | ".join(c.text.strip() for c in row.cells) for row in t.rows]
            text = "\n".join(parts)
        elif suffix in (".xlsx", ".xlsm"):
            import openpyxl

            values = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
            formulas = openpyxl.load_workbook(io.BytesIO(content), data_only=False, read_only=True)
            parts = []
            for ws, wf in zip(values.worksheets, formulas.worksheets):
                parts.append(f"## sheet {ws.title!r}")
                for rv, rf in zip(ws.iter_rows(values_only=True), wf.iter_rows(values_only=True)):
                    # A formula openpyxl wrote has no cached value: show the formula.
                    cells = ["" if (v if v is not None else f) is None else str(v if v is not None else f)
                             for v, f in zip(rv, rf)]
                    if any(cells):
                        parts.append(" | ".join(cells).rstrip(" |"))
            text = "\n".join(parts)
        elif suffix == ".pptx":
            import pptx

            p = pptx.Presentation(io.BytesIO(content))
            parts = []
            for i, slide in enumerate(p.slides, 1):
                parts.append(f"## slide {i}")
                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text_frame.text.strip():
                        parts.append(shape.text_frame.text)
                    if getattr(shape, "has_table", False) and shape.has_table:
                        parts += [" | ".join(c.text for c in row.cells) for row in shape.table.rows]
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                    parts.append(f"(notes) {slide.notes_slide.notes_text_frame.text}")
            text = "\n".join(parts)
        elif suffix == ".pdf":
            import pypdf

            reader = pypdf.PdfReader(io.BytesIO(content))
            text = "\n".join(f"## page {i}\n{page.extract_text() or ''}" for i, page in enumerate(reader.pages, 1))
        elif suffix == ".zip":
            parts = []
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                for info in z.infolist():
                    if info.is_dir():
                        continue
                    parts.append(f"### {info.filename}\n" + render_file(info.filename, z.read(info), limit // 4))
            text = "\n".join(parts)
        elif suffix in TEXT_SUFFIXES or not suffix:
            text = content.decode("utf-8")
        else:
            return f"[binary {suffix or 'file'}, {len(content)} bytes: not shown]"
    except Exception as e:
        return f"[{suffix or 'file'}, {len(content)} bytes: could not be read ({type(e).__name__}: {e})]"
    if len(text) > limit:
        text = text[:limit] + f"\n[... truncated, {len(text) - limit} more characters]"
    return text


def render_submission(files: list[tuple[str, bytes]], note: str | None = None, limit: int = 120_000) -> str:
    parts = []
    if note:
        parts.append(f"(message accompanying the submission) {note}")
    if not files:
        parts.append("(no files were submitted)")
    per_file = max(limit // max(len(files), 1), 4000)
    for name, content in files:
        parts.append(f"--- file: {name} ({len(content)} bytes) ---\n{render_file(name, content, per_file)}")
    return "\n\n".join(parts)


def renderable(name: str) -> bool:
    suffix = Path(name).suffix.lower()
    return suffix in {".docx", ".xlsx", ".xlsm", ".pptx", ".pdf", ".zip"} | TEXT_SUFFIXES


@functools.cache
def sandbox_provider_class():
    """Stirrup code-exec provider over a granite :class:`~granite_evals.sandbox.Sandbox`
    (enroot). Built lazily: stirrup is imported only with the extra."""
    from stirrup.tools.code_backends.base import CodeExecToolProvider, CommandResult, OutputSourceRoot

    from granite_evals.sandbox import make_sandbox

    class SandboxCodeExecToolProvider(CodeExecToolProvider):
        def __init__(self, image: str, *, backend: str = "", workdir: str = "/workspace",
                     setup: str = "", setup_timeout: int = 1800, shell_timeout: int = 600, factory=None):
            super().__init__(shell_timeout=shell_timeout)
            self.image, self.backend, self.workdir = image, backend, workdir
            self.setup, self.setup_timeout = setup, setup_timeout
            self._factory = factory or (lambda: make_sandbox(image, backend=backend))
            self._sb = None
            self.setup_output = ""

        def _abs(self, path: str) -> str:
            return path if path.startswith("/") else f"{self.workdir}/{path}"

        def _exec(self, cmd: str, timeout: int, stdin: str | None = None):
            return self._sb.execute(cmd, cwd=self.workdir, timeout=timeout, stdin=stdin)

        async def __aenter__(self):
            sb = self._factory()
            if hasattr(sb, "_run_env"):
                # enroot hands the caller's environment to the container: keep keys out of
                # the sandbox, where model-written code runs.
                inner = sb._run_env

                def scrubbed():
                    env = inner() or dict(os.environ)
                    return {k: v for k, v in env.items() if not _SECRET_ENV.search(k)}

                sb._run_env = scrubbed
            await asyncio.to_thread(sb.start)
            self._sb = sb
            r = await asyncio.to_thread(sb.execute, f"mkdir -p {shlex.quote(self.workdir)}", cwd="/", timeout=120)
            if r.returncode != 0:
                raise RuntimeError(f"sandbox workdir: {r.output[-2000:]}")
            if self.setup:
                r = await asyncio.to_thread(self._exec, self.setup, self.setup_timeout)
                self.setup_output = r.output[-4000:]
                if r.returncode != 0:
                    raise RuntimeError(f"sandbox setup failed (exit {r.returncode}): {r.output[-2000:]}")
            return self.get_code_exec_tool()

        async def __aexit__(self, exc_type, exc_val, exc_tb):
            if self._sb is not None:
                await asyncio.to_thread(self._sb.close)
                self._sb = None

        async def run_command(self, cmd: str, *, timeout: int | None = None):
            t = timeout or self._shell_timeout
            wrapped = f"timeout --kill-after=5s {t}s bash -c {shlex.quote(cmd)}"
            r = await asyncio.to_thread(self._exec, wrapped, t + 60)
            if r.timed_out or r.returncode in (124, 137):
                return CommandResult(exit_code=r.returncode, stdout=r.output, stderr=f"command timed out after {t}s",
                                     error_kind="timeout", advice="Use a shorter command or run it in the background.")
            return CommandResult(exit_code=r.returncode, stdout=r.output, stderr="")

        async def read_file_bytes(self, path: str) -> bytes:
            q = shlex.quote(self._abs(path))
            r = await asyncio.to_thread(self._exec, f"test -f {q} && base64 < {q} | tr -d '\\n'", 300)
            if r.returncode != 0:
                raise FileNotFoundError(path)
            return base64.b64decode(r.output.strip())

        async def write_file_bytes(self, path: str, content: bytes) -> None:
            p = self._abs(path)
            cmd = f"mkdir -p {shlex.quote(os.path.dirname(p))} && base64 -d > {shlex.quote(p)}"
            r = await asyncio.to_thread(self._exec, cmd, 300, base64.b64encode(content).decode())
            if r.returncode != 0:
                raise RuntimeError(f"writing {path}: {r.output[-1000:]}")

        async def file_exists(self, path: str) -> bool:
            r = await asyncio.to_thread(self._exec, f"test -f {shlex.quote(self._abs(path))}", 60)
            return r.returncode == 0

        async def is_directory(self, path: str) -> bool:
            r = await asyncio.to_thread(self._exec, f"test -d {shlex.quote(self._abs(path))}", 60)
            return r.returncode == 0

        async def list_files(self, path: str) -> list[str]:
            q = shlex.quote(self._abs(path))
            r = await asyncio.to_thread(self._exec, f"test -d {q} && cd {q} && find . -type f | sed 's|^[.]/||'", 120)
            return [line for line in r.output.splitlines() if line] if r.returncode == 0 else []

        def output_source_roots(self):
            return (OutputSourceRoot(path=self.workdir),)

        async def view_image(self, path: str):
            raise NotImplementedError("no image viewer in this harness")

    return SandboxCodeExecToolProvider


@functools.cache
def quiet_logger_class():
    """A Stirrup AgentLoggerBase that logs through ``logging`` (the default one
    takes over the root logger with a rich console)."""
    from stirrup.utils.logging import AgentLoggerBase

    class QuietAgentLogger(AgentLoggerBase):
        def __init__(self, label: str):
            self.label = label
            self.name, self.model, self.max_turns, self.depth = label, None, None, 0
            self.finish_params = self.run_metadata = self.output_dir = None

        def __enter__(self):
            return self

        def __exit__(self, *exc):
            pass

        def on_step(self, step, tool_calls=0, input_tokens=0, output_tokens=0):
            log.info("gdpval: %s step %d tool_calls=%d tokens_in=%d tokens_out=%d",
                     self.label, step, tool_calls, input_tokens, output_tokens)

        def assistant_message(self, turn, max_turns, assistant_message):
            pass

        def user_message(self, user_message):
            pass

        def task_message(self, task):
            pass

        def tool_result(self, tool_message):
            pass

        def context_summarization_start(self, pct_used, cutoff):
            log.info("gdpval: %s summarizing context at %.0f%%", self.label, 100 * pct_used)

        def context_summarization_complete(self, summary, bridge):
            pass

        def debug(self, message, *args):
            log.debug(f"gdpval: {self.label} {message}", *args)

        def info(self, message, *args):
            log.debug(f"gdpval: {self.label} {message}", *args)

        def warning(self, message, *args):
            log.warning(f"gdpval: {self.label} {message}", *args)

        def error(self, message, *args):
            log.error(f"gdpval: {self.label} {message}", *args)

    return QuietAgentLogger


def _dump_history(history) -> list:
    out = []
    for group in history or []:
        out.append([m.model_dump(mode="json") if hasattr(m, "model_dump") else str(m) for m in group])
    return out


@register
class GDPval(JudgedBenchmark):
    """GDPval gold set: Stirrup agent in a sandbox, pairwise-judged against the expert.

    Options:
      deliverables=model|expert   grade the model's work (default) or the expert's own
                                  deliverables against themselves (pipeline check: ~0.5)
      sandbox_image=python:3.13-bookworm   workspace image (any pullable image with python3)
      sandbox_setup=<shell>       run once per task before the agent (default: pip install
                                  the office-file libraries); empty to skip
      sandbox=enroot|podman|docker
      max_turns=250  shell_timeout=600  max_tokens=32768  context_window=131072
      temperature, top_p          policy sampling (default: generation_config)
      judge_model / judge_base_url / judge_api_key_env, judge_max_tokens=16384,
      judge_reasoning_effort=     (unset = provider default)
      judge_thinking, judge_effort, judge_extra_body
                                  provider thinking fields (JudgedBenchmark.judge_extra_body)
      judge_orders=2             1 = model as A only; 2 = both orders
      elo_anchor=1290             Elo-style value of the expert
      max_failed_frac=0.05        fail the run above this fraction of failed tasks (0 = any)

    A task whose agent, sandbox or judge call raised (error:*), or whose judge
    gave no valid verdict in any order (judge_invalid), is a failure, not a
    loss: it is left out of the Elo and win rates and retried on resume (the
    agent's saved run is reused, so only the judging is redone). GDPval ships
    no automated grader to follow here; this matches the other judged
    benchmarks. The run reports tasks_failed / tasks_total over the tasks not
    excluded, fails above max_failed_frac and is flagged incomplete otherwise.
    A task the model gave up on (no_submission) is still a loss.

    ``--phase generate`` runs the agents (``tasks/<id>/agent.json`` and
    ``deliverables/``) and writes no report.json; ``--phase score`` judges them.
    A task without agent.json in the score phase is ``error:not_generated``, a
    failure like error:* (left out, retried on resume). Exclusions are decided,
    and their reports written, by the scoring run, so n and the metrics match
    ``--phase all``.
    """

    id = "gdpval"
    metric = "Elo"
    harness_packages = ("stirrup", "openai", "python-docx", "openpyxl", "python-pptx", "pypdf")
    dataset = "openai/gdpval"
    dataset_revision = "11e7900cdcac61bc4daf59e65feb238acda98fbf"
    split: ClassVar[str] = "train"
    splittable = True
    default_setup: ClassVar[str] = (
        "python3 -m pip install --quiet --disable-pip-version-check --root-user-action=ignore "
        "python-docx openpyxl python-pptx reportlab pandas matplotlib xlsxwriter pypdf >/tmp/setup.log 2>&1 "
        "|| (tail -20 /tmp/setup.log; exit 1)"
    )

    @property
    def expert_mode(self) -> bool:
        return self.opt("deliverables", "model") == "expert"

    def needs_server(self) -> bool:
        return not self.expert_mode

    def score_needs_server(self) -> bool:
        return self.judge_model == "self"

    # -- entry point -------------------------------------------------------

    def run(self, base_url: str, served_model_name: str) -> dict[str, Any]:
        self.one_sample()
        source, revision = self.dataset_source()
        rows = data.load_split(source, revision=revision, split=self.split)
        if pattern := self.opt("tasks", ""):
            rows = [r for r in rows if re.search(pattern, r["task_id"])]
        tasks = data.take(rows, self.config.limit, key="task_id")
        root = self.config.output_dir / "tasks"
        if not self.scoring:
            log.info("gdpval: %d tasks, generate only (deliverables=%s)", len(tasks),
                     "expert" if self.expert_mode else "model")
            with concurrent.futures.ThreadPoolExecutor(max_workers=self.config.workers) as pool:
                runs = list(pool.map(lambda t: self._generate_task(t, root, base_url, served_model_name,
                                                                   source, revision), tasks))
            return self._generated(runs, tasks, source, revision)
        judge = self.make_judge(base_url, served_model_name, _gdp_cache_split)
        log.info("gdpval: %d tasks (deliverables=%s, judge=%s)", len(tasks),
                 "expert" if self.expert_mode else "model", judge.model)
        with concurrent.futures.ThreadPoolExecutor(max_workers=self.config.workers) as pool:
            reports = list(pool.map(lambda t: self._task(t, root, base_url, served_model_name, judge, source, revision), tasks))
        return self._aggregate(reports, tasks, judge, source, revision)

    # -- files -------------------------------------------------------------

    def _fetch(self, source: str, revision: str | None, rel: str) -> Path:
        if Path(source).exists():
            return Path(source) / rel
        from huggingface_hub import hf_hub_download

        return Path(hf_hub_download(source, rel, repo_type="dataset", revision=revision))

    # -- one task ----------------------------------------------------------

    def _task(self, task: dict, root: Path, base_url: str, served: str, judge: Judge, source, revision) -> dict:
        tid = task["task_id"]
        tdir = root / tid
        tdir.mkdir(parents=True, exist_ok=True)
        report_path = tdir / "report.json"
        if report_path.exists():
            report = json.loads(report_path.read_text())
            if report["status"] != "judge_invalid":  # judge_invalid is re-judged
                return report
        base = {"task_id": tid, "sector": task["sector"], "occupation": task["occupation"]}
        try:
            expert = [(Path(p).name, self._fetch(source, revision, p).read_bytes()) for p in task["deliverable_files"]]
            if not expert:
                report = {**base, "status": "excluded:no_expert_deliverable"}
            elif not any(renderable(name) for name, _ in expert):
                report = {**base, "status": "excluded:expert_deliverable_not_text"}
            else:
                candidate = self._candidate(task, tdir, [n for n, _ in expert], base_url, served, source, revision)
                if candidate["status"].startswith("error:"):
                    log.warning("gdpval: %s", candidate["error"])
                    return {**base, **candidate}  # not generated: retried on resume, no report
                if candidate["status"] != "submitted":
                    report = {**base, **candidate, "score": 0.0}  # nothing to grade: a loss
                else:
                    report = {**base, **candidate, **self._judge(task, expert, candidate, judge, tdir)}
        except Exception as e:  # one broken task must not sink the run
            log.exception("gdpval: %s failed", tid)
            return {**base, "status": f"error:{type(e).__name__}", "error": str(e)[:500]}  # retried on resume
        _write_json(report_path, report)
        log.info("gdpval: %s %s score=%s", tid, report["status"], report.get("score"))
        return report

    def _generate_task(self, task: dict, root: Path, base_url: str, served: str, source, revision) -> dict:
        """Generate phase: the agent's run for one task, saved, not judged.

        Exclusion is read from the expert file names (not fetched); the scoring
        run decides it again from the files and writes the report."""
        tid = task["task_id"]
        tdir = root / tid
        tdir.mkdir(parents=True, exist_ok=True)
        base = {"task_id": tid}
        names = [Path(p).name for p in task["deliverable_files"]]
        if not names:
            return {**base, "status": "excluded:no_expert_deliverable"}
        if not any(renderable(n) for n in names):
            return {**base, "status": "excluded:expert_deliverable_not_text"}
        try:
            candidate = self._candidate(task, tdir, names, base_url, served, source, revision)
        except Exception as e:  # one broken task must not sink the run
            log.exception("gdpval: %s failed", tid)
            return {**base, "status": f"error:{type(e).__name__}", "error": str(e)[:500]}  # retried on resume
        log.info("gdpval: %s generated: %s", tid, candidate["status"])
        return {**base, "status": candidate["status"]}

    def _candidate(self, task, tdir, expert_names, base_url, served, source, revision) -> dict:
        """The submission to grade: {status, files: [names], note, agent: {...}}."""
        out_dir = tdir / "deliverables"
        agent_path = tdir / "agent.json"
        if self.expert_mode:
            return {"status": "submitted", "files": list(expert_names), "note": None, "candidate": "expert"}
        if not agent_path.exists():
            if not self.generating:  # score phase: never run the agent
                return {"status": "error:not_generated", "error": str(self.not_generated(f"task {task['task_id']}"))}
            refs = [self._fetch(source, revision, p) for p in task["reference_files"]]
            _write_json(agent_path, asyncio.run(self._agent(task, refs, out_dir, base_url, served)))
        agent = json.loads(agent_path.read_text())
        files = sorted(str(p.relative_to(out_dir)) for p in out_dir.rglob("*") if p.is_file()) if out_dir.exists() else []
        status = "submitted" if agent["finished"] and (files or (agent.get("note") or "").strip()) else "no_submission"
        return {"status": status, "files": files, "note": agent.get("note"), "candidate": "model",
                "agent": {k: agent[k] for k in ("finished", "turns", "token_usage", "failed_outputs") if k in agent}}

    async def _agent(self, task: dict, refs: list[Path], out_dir: Path, base_url: str, served: str) -> dict:
        from stirrup import Agent
        from stirrup.clients.chat_completions_client import ChatCompletionsClient

        client = ChatCompletionsClient(
            model=served,
            max_tokens=self.opt("max_tokens", 32768),
            context_window_tokens=self.opt("context_window", 131072),
            base_url=base_url,
            api_key="EMPTY",
            timeout=self.opt("timeout", 3600.0),
            kwargs={"seed": self.config.seed, **self.sampling()},
        )
        provider = sandbox_provider_class()(
            self.opt("sandbox_image", "python:3.13-bookworm"),
            backend=self.opt("sandbox", ""),
            setup=self.opt("sandbox_setup", self.default_setup),
            shell_timeout=self.opt("shell_timeout", 600),
        )
        agent = Agent(
            client=client,
            name="gdpval",
            max_turns=self.opt("max_turns", 250),
            system_prompt=GDPVAL_INSTRUCTIONS,
            tools=[provider],
            logger=quiet_logger_class()(task["task_id"][:8]),
        )
        async with agent.session(output_dir=str(out_dir), input_files=[str(p) for p in refs] or None,
                                 cache_on_interrupt=False, clear_cache_on_success=False) as session:
            finish, history, metadata = await session.run(task["prompt"])
        (out_dir.parent / "history.json").write_text(json.dumps(_dump_history(history), default=str))
        usage = metadata.get("token_usage") or []
        paths = list(getattr(finish, "paths", []) or [])
        return {
            "finished": finish is not None,
            "note": getattr(finish, "reason", None),
            "paths": paths,
            "turns": sum(1 for g in history for m in g if type(m).__name__ == "AssistantMessage"),
            "token_usage": {
                "input": sum(u.input for u in usage),
                "output": sum(u.answer + u.reasoning for u in usage),
                "reasoning": sum(u.reasoning for u in usage),
            },
            # Reported deliverables that did not arrive in out_dir.
            "failed_outputs": [p for p in paths if not (out_dir / p.removeprefix("/workspace/").lstrip("/")).exists()],
            "setup_output": provider.setup_output,
        }

    # -- judging -------------------------------------------------------------

    def _judge(self, task, expert, candidate, judge: Judge, tdir: Path) -> dict:
        if candidate["candidate"] == "expert":
            cand_files, note = expert, None
        else:
            out_dir = tdir / "deliverables"
            cand_files = [(n, (out_dir / n).read_bytes()) for n in candidate["files"]]
            note = candidate["note"]
        model_text = render_submission(cand_files, note)
        expert_text = render_submission(expert)
        refs = "\n".join(f"- {Path(p).name}" for p in task["reference_files"]) or "(none)"
        orders = self.opt("judge_orders", 2)
        verdicts, usage, raw = [], [], []
        for model_is_a in (True, False)[:orders]:
            prompt = GDPVAL_JUDGE_PROMPT.format(
                occupation=task["occupation"], sector=task["sector"], prompt=task["prompt"],
                references=refs, rubric=task.get("rubric_pretty") or "(none)",
                a=model_text if model_is_a else expert_text, b=expert_text if model_is_a else model_text,
            )
            request = {"model": judge.model, "messages": [{"role": "user", "content": prompt}],
                       "max_tokens": self.opt("judge_max_tokens", 16384)}
            if effort := self.opt("judge_reasoning_effort", ""):
                request["reasoning_effort"] = effort
            completion = judge.chat.completions.create(**request)
            text = completion.choices[0].message.content or ""
            usage.append(judge.last_usage())
            v = parse_verdict(text)
            raw.append(text)
            if v is None:
                verdicts.append(None)
            else:
                model_side = "A" if model_is_a else "B"
                verdicts.append(0.5 if v == "TIE" else float(v == model_side))
        (tdir / "judge.json").write_text(json.dumps({"responses": raw, "verdicts": verdicts}, indent=2))
        valid = [v for v in verdicts if v is not None]
        if not valid:
            return {"status": "judge_invalid", "verdicts": verdicts, "judge_usage": usage}
        return {"status": "judged", "score": sum(valid) / len(valid), "verdicts": verdicts, "judge_usage": usage}

    # -- aggregate -----------------------------------------------------------

    def _generated(self, runs, tasks, source, revision) -> dict[str, Any]:
        """The generate phase's summary (generation.json): agents run, nothing judged."""
        failed = sum(r["status"].startswith("error:") for r in runs)
        if failed:
            log.warning("gdpval: %d/%d agent runs failed; rerun --phase generate to retry them", failed, len(runs))
        return {
            "n": sum(r["status"] in ("submitted", "no_submission") for r in runs),  # tasks with a saved run
            "tasks_failed": failed,
            "statuses": _count(r["status"] for r in runs),
            "dataset": source,
            "dataset_revision": revision,
            "deliverables": "expert" if self.expert_mode else "model",
            "tasks": [t["task_id"] for t in tasks],
            "harness": "stirrup (ArtificialAnalysis/Stirrup), code_exec + finish tools, granite sandbox",
            "sandbox_image": self.opt("sandbox_image", "python:3.13-bookworm"),
            "sampling": {**self.sampling(), "max_tokens": self.opt("max_tokens", 32768),
                         "defaults": "generation_config of the checkpoint"},
        }

    def _aggregate(self, reports, tasks, judge: Judge, source, revision) -> dict[str, Any]:
        graded = [r for r in reports if r["status"] in ("judged", "no_submission")]
        n = len(graded)
        failed = sum(r["status"] == "judge_invalid" or r["status"].startswith("error:") for r in reports)
        total = sum(not r["status"].startswith("excluded:") for r in reports)
        policy = failure_policy(self.id, "tasks", failed, total, self.opt("max_failed_frac", 0.05))
        if failed:
            log.warning("gdpval: %d/%d tasks failed (agent, sandbox or judge), left out; rerun to retry them",
                        failed, total)
        if total and not n:
            raise SystemExit("gdpval: no task was graded")
        score = sum(r["score"] for r in graded)
        anchor = self.opt("elo_anchor", 1290.0)
        usage = [u for r in reports for u in r.get("judge_usage", [])]
        judged = [r for r in reports if r["status"] == "judged"]
        elo = elo_style(score, n, anchor)
        note = GDPVAL_METRIC_NOTE.format(judge=judge.model + (" (the served model)" if judge.is_self else ""), anchor=anchor)
        out = {
            "value": round(elo, 1),
            "n": n,
            "dataset": source,
            "dataset_revision": revision,
            "elo_is_approximation": True,
            "metric_note": note,
            "elo_anchor": anchor,
            "win_rate": score / n if n else 0.0,
            "wins_or_ties_rate": sum(r["score"] >= 0.5 for r in graded) / n if n else 0.0,
            "pass_at_k": pass_at_k_record(1, score / n if n else 0.0, None, n,
                                          "pass_at_1 is the win rate the Elo comes from; no pass@k for an Elo"),  # fmt: skip
            "deliverables": "expert" if self.expert_mode else "model",
            "statuses": _count(r["status"] for r in reports),
            **policy,
            "tasks": [t["task_id"] for t in tasks],
            "per_sector": {
                s: round(sum(r["score"] for r in graded if r["sector"] == s)
                         / max(1, sum(r["sector"] == s for r in graded)), 4)
                for s in sorted({r["sector"] for r in graded})
            },
            "position_consistency": (
                sum(len(set(r["verdicts"])) == 1 for r in judged if None not in r["verdicts"]) / len(judged)
                if judged else None
            ),
            "harness": "stirrup (ArtificialAnalysis/Stirrup), code_exec + finish tools, granite sandbox",
            "sandbox_image": self.opt("sandbox_image", "python:3.13-bookworm"),
            "sampling": {**self.sampling(), "max_tokens": self.opt("max_tokens", 32768),
                         "defaults": "generation_config of the checkpoint"},
            **judge.describe(),
            "judge_usage": total_usage(usage, len(judged)),
        }
        log.warning("gdpval: Elo-style %.1f from win rate %.3f vs the expert over %d tasks -- %s",
                    elo, out["win_rate"], n, note)
        return out
